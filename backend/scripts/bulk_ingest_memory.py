import os
import sys
import asyncio
import uuid
import re

# Add the backend directory to python path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from services.vector_service import VectorSearchService
from services.llm_service import LLMService

KNOWLEDGE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../knowledge_base/extended learning on LOBs'))

def extract_text_from_docx(file_path):
    import docx
    try:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    except Exception as e:
        print(f"Error reading docx {file_path}: {e}")
        return ""

def extract_text_from_pptx(file_path):
    from pptx import Presentation
    try:
        prs = Presentation(file_path)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)
    except Exception as e:
        print(f"Error reading pptx {file_path}: {e}")
        return ""

def chunk_text(text, max_words=300):
    words = text.split()
    chunks = []
    for i in range(0, len(words), max_words):
        chunks.append(" ".join(words[i:i + max_words]))
    return chunks

async def process_files():
    if not os.path.exists(KNOWLEDGE_DIR):
        print(f"Directory not found: {KNOWLEDGE_DIR}")
        return

    vector_service = VectorSearchService(index_name="insurance-guidelines")
    llm_service = LLMService()

    for root, dirs, files in os.walk(KNOWLEDGE_DIR):
        for file in files:
            if file.startswith("~"): continue # Skip temp files
            
            file_path = os.path.join(root, file)
            rel_path = os.path.relpath(root, KNOWLEDGE_DIR)
            
            # e.g. "Personal Lines\Auto" -> parts = ["Personal Lines", "Auto"]
            parts = rel_path.split(os.sep)
            category = parts[0] if len(parts) > 0 and parts[0] != '.' else "General"
            lob = parts[1] if len(parts) > 1 else "Unknown"

            print(f"Processing: {file} (Category: {category}, LOB: {lob})")
            
            text = ""
            if file.endswith(".docx"):
                text = extract_text_from_docx(file_path)
            elif file.endswith(".pptx"):
                text = extract_text_from_pptx(file_path)
            else:
                print(f"Unsupported file type: {file}")
                continue

            if not text.strip():
                print(f"No text extracted from {file}")
                continue

            chunks = chunk_text(text)
            print(f" -> Generated {len(chunks)} chunks.")

            for i, chunk in enumerate(chunks):
                # Clean up chunk to avoid weird encoding issues
                clean_chunk = re.sub(r'[^\x00-\x7F]+', ' ', chunk)
                
                print(f"   -> Generating embedding for chunk {i+1}/{len(chunks)}...")
                vector = await llm_service.get_embeddings(clean_chunk)
                
                doc_id = str(uuid.uuid4())
                
                # Push to Azure AI Search
                document = {
                    "id": doc_id,
                    "content": clean_chunk,
                    "requirement_id": "", # Not applicable for guidelines
                    "doc_id": file,
                    "lob": lob,
                    "category": category, # New field
                    "content_vector": vector,
                    "metadata": str({"chunk_index": i})
                }
                
                try:
                    vector_service.search_client.upload_documents(documents=[document])
                except Exception as e:
                    print(f"   -> ERROR uploading chunk: {e}")

    print(" Bulk ingestion complete!")

if __name__ == "__main__":
    asyncio.run(process_files())
