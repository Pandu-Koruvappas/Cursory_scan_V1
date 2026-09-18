import os
import uuid
import zipfile
import fitz
import docx
import openpyxl
import pptx
import aiofiles

async def extract_all_text(file_path: str, filename: str, max_depth: int = 3, current_depth: int = 0) -> str:
    if current_depth > max_depth:
        return f"\n[Max extraction depth reached for {filename}]\n"
        
    print(f"{'  '*current_depth}[INFO] Deep parsing: {filename}")
    ext = os.path.splitext(filename)[1].lower()
    text = f"\n--- BEGIN DOCUMENT: {filename} ---\n"
    
    try:
        if ext == ".pdf":
            text += await _parse_pdf(file_path, current_depth, max_depth)
        elif ext == ".docx":
            text += await _parse_docx(file_path, current_depth, max_depth)
        elif ext == ".xlsx":
            text += await _parse_xlsx(file_path, current_depth, max_depth)
        elif ext == ".pptx":
            text += await _parse_pptx(file_path, current_depth, max_depth)
        else:
            async with aiofiles.open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text += await f.read()
    except Exception as e:
        text += f"\n[Error extracting {filename}: {str(e)}]\n"
        
    text += f"\n--- END DOCUMENT: {filename} ---\n"
    return text

async def _parse_pdf(file_path: str, depth: int, max_depth: int) -> str:
    fitz_text = ""
    total_pages = 0
    try:
        doc = fitz.open(file_path)
        total_pages = len(doc)
        for i, page in enumerate(doc):
            p_text = page.get_text()
            if p_text and p_text.strip():
                fitz_text += f"\n--- Page {i+1} ---\n" + p_text.strip() + "\n"
        doc.close()
    except Exception as err:
        print(f"{'  '*depth}[WARN] PyMuPDF extraction error for {file_path}: {err}")

    adi_text = ""
    try:
        from services.adi_service import AzureDocIntelService
        adi = AzureDocIntelService()
        adi_res = adi.extract_text(file_path)
        if adi_res and isinstance(adi_res, str):
            adi_text = adi_res
    except Exception as e:
        print(f"{'  '*depth}[WARN] ADI extraction skipped/failed for {file_path}: {e}")

    # Select fullest text extraction
    if len(adi_text.strip()) > len(fitz_text.strip()):
        final_text = adi_text
    else:
        final_text = fitz_text

    print(f"{'  '*depth}[SUCCESS] Extracted {len(final_text)} characters across {total_pages} pages from {os.path.basename(file_path)}")

    # Extract embedded files using PyMuPDF
    try:
        doc = fitz.open(file_path)
        emb_names = doc.embfile_names()
        for name in emb_names:
            print(f"{'  '*depth}[EMBED] Found embedded file in PDF: {name}")
            data = doc.embfile_get(name)
            temp_emb_path = f"temp_emb_{uuid.uuid4()}_{name}"
            async with aiofiles.open(temp_emb_path, "wb") as f:
                await f.write(data)
            
            emb_text = await extract_all_text(temp_emb_path, name, max_depth, depth + 1)
            final_text += f"\n\n{emb_text}\n"
            
            if os.path.exists(temp_emb_path):
                os.remove(temp_emb_path)
        doc.close()
    except Exception as e:
        print(f"{'  '*depth}[ERROR] Error extracting PDF attachments: {e}")
        
    return final_text

async def _parse_pptx(file_path: str, depth: int, max_depth: int) -> str:
    text = ""
    try:
        prs = pptx.Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            text += f"\n--- Slide {i+1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
                if shape.has_table:
                    for row in shape.table.rows:
                        row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_data:
                            text += " | ".join(row_data) + "\n"
    except Exception as e:
        text += f"\n[Error parsing PPTX text: {str(e)}]\n"

    # Extract embedded OLE objects
    try:
        with zipfile.ZipFile(file_path, "r") as z:
            for item in z.namelist():
                if item.startswith("ppt/embeddings/") and "." in item:
                    basename = os.path.basename(item)
                    if not basename: continue
                    
                    data = z.read(item)
                    temp_emb_path = f"temp_emb_{uuid.uuid4()}_{basename}"
                    async with aiofiles.open(temp_emb_path, "wb") as f:
                        await f.write(data)
                        
                    emb_text = await extract_all_text(temp_emb_path, basename, max_depth, depth + 1)
                    text += f"\n\n[Embedded File: {basename}]\n{emb_text}\n"
                    
                    if os.path.exists(temp_emb_path):
                        os.remove(temp_emb_path)
    except Exception as e:
        print(f"{'  '*depth}[ERROR] Error extracting PPTX attachments: {e}")

    return text

async def _parse_docx(file_path: str, depth: int, max_depth: int) -> str:
    text = ""
    
    # 1. Parse main text
    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            if p.text.strip():
                text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    text += " | ".join(row_data) + "\n"
    except Exception as e:
        text += f"\n[Error parsing DOCX text: {str(e)}]\n"

    # 2. Extract embedded OLE objects (Excel/Word)
    try:
        with zipfile.ZipFile(file_path, "r") as z:
            for item in z.namelist():
                # Word embeds files in word/embeddings/ (e.g. oleObject1.bin or Microsoft_Excel_Worksheet1.xlsx)
                if item.startswith("word/embeddings/") and "." in item:
                    # Map common bin extensions if needed, or extract directly
                    # Sometimes they are raw .bin OLE streams, which are harder to parse.
                    # But often modern Office embeds them as .xlsx or .docx
                    basename = os.path.basename(item)
                    if not basename: continue
                    
                    # Microsoft OLE files embedded as .bin usually have standard office files inside or we can attempt to rename them if we know what they are. 
                    # If they are oleObject1.bin, openpyxl/docx will likely fail to read them natively unless we use an OLE parser like olefile, but we'll try as is.
                    print(f"{'  '*depth}[EMBED] Found embedded file in DOCX: {basename}")
                    
                    data = z.read(item)
                    
                    # If it's a generic .bin and we think it's excel, we might change the extension, but let's stick to the name inside the zip for now
                    temp_emb_path = f"temp_emb_{uuid.uuid4()}_{basename}"
                    async with aiofiles.open(temp_emb_path, "wb") as f:
                        await f.write(data)
                    
                    emb_text = await extract_all_text(temp_emb_path, basename, max_depth, depth + 1)
                    text += f"\n\n{emb_text}\n"
                    
                    if os.path.exists(temp_emb_path):
                        os.remove(temp_emb_path)
    except zipfile.BadZipFile:
        pass
    except Exception as e:
        print(f"{'  '*depth}[ERROR] Error extracting DOCX attachments: {e}")
        
    return text

async def _parse_xlsx(file_path: str, depth: int, max_depth: int) -> str:
    text = ""
    
    # 1. Parse main text
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text += f"\nSheet: {sheet_name}\n"
            for row in sheet.iter_rows(values_only=True):
                row_data = [str(v) for v in row if v is not None and str(v).strip()]
                if row_data:
                    text += " | ".join(row_data) + "\n"
        wb.close()
    except Exception as e:
        text += f"\n[Error parsing XLSX text: {str(e)}]\n"

    # 2. Extract embedded OLE objects
    try:
        with zipfile.ZipFile(file_path, "r") as z:
            for item in z.namelist():
                if item.startswith("xl/embeddings/") and "." in item:
                    basename = os.path.basename(item)
                    if not basename: continue
                    print(f"{'  '*depth}[EMBED] Found embedded file in XLSX: {basename}")
                    
                    data = z.read(item)
                    temp_emb_path = f"temp_emb_{uuid.uuid4()}_{basename}"
                    async with aiofiles.open(temp_emb_path, "wb") as f:
                        await f.write(data)
                        
                    emb_text = await extract_all_text(temp_emb_path, basename, max_depth, depth + 1)
                    text += f"\n\n{emb_text}\n"
                    
                    if os.path.exists(temp_emb_path):
                        os.remove(temp_emb_path)
    except zipfile.BadZipFile:
        pass
    except Exception as e:
        print(f"{'  '*depth}[ERROR] Error extracting XLSX attachments: {e}")
        
    return text
